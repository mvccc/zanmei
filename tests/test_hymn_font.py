from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from mvccc.slides import FLAGS, Hymn, hymn_font_size


def test_crowded_page_controls_song_size():
    short = [["短句"] * 3]
    crowded = [["較長的中文歌詞"] * 5]
    assert hymn_font_size(short + crowded, 880, 400) == hymn_font_size(crowded, 880, 400)
    assert hymn_font_size(short + crowded, 880, 400) < hymn_font_size(short, 880, 400)
    assert hymn_font_size([["長" * 30]], 880, 400) < hymn_font_size(short, 880, 400)


def test_generated_song_has_one_size_and_no_autofit(tmp_path):
    if not FLAGS.is_parsed():
        FLAGS(["test"])
    template = Path(__file__).resolve().parents[1] / "mvccc_master_keynote_blue.pptx"
    deck = Presentation(template)
    start = len(deck.slides)
    Hymn("test", [("", ["歌詞"] * n) for n in (3, 4, 5)]).add_to(deck)
    output = tmp_path / "song.pptx"
    deck.save(output)
    deck = Presentation(output)
    sizes = set()
    for slide in list(deck.slides)[start + 1 :]:
        body = max((s for s in slide.shapes if s.has_text_frame), key=lambda s: len(s.text))
        assert body.text_frame.auto_size == MSO_AUTO_SIZE.NONE
        assert body.text_frame.word_wrap is False
        sizes.update(run.font.size for p in body.text_frame.paragraphs for run in p.runs)
    assert len(sizes) == 1
