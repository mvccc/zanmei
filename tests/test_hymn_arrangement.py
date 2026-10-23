from pathlib import Path

import pytest
from pptx import Presentation

from mvccc.slides import FLAGS, Hymn, search_hymn_md


def test_repeat_reorder_and_preserve_omitted_pages(tmp_path):
    if not FLAGS.is_parsed():
        FLAGS(["test"])
    template = Path(__file__).resolve().parents[1] / "mvccc_master_keynote_blue.pptx"
    deck = Presentation(template)
    start = len(deck.slides)
    Hymn("test", [("", [text]) for text in ("first", "second", "third")], page_order=[2, 1, 2]).add_to(deck)
    path = tmp_path / "arranged.pptx"
    deck.save(path)
    slides = list(Presentation(path).slides)[start + 1 :]
    assert [max((s.text for s in slide.shapes if s.has_text_frame), key=len) for slide in slides] == [
        "second",
        "first",
        "second",
        "third",
    ]
    assert [slide._element.get("show", "1") for slide in slides] == ["1", "1", "1", "0"]


def test_page_suffix_and_default_order(tmp_path):
    (tmp_path / "001_song.md").write_text("# Song\n\n##\none\n\n##\ntwo\n")
    assert search_hymn_md("001_song", tmp_path)[0].page_order is None
    assert search_hymn_md("001_song@2,1,2", tmp_path)[0].page_order == [2, 1, 2]
    for suffix in ("", "0", "-1", "1,", "abc", "3"):
        with pytest.raises(ValueError):
            search_hymn_md("001_song@" + suffix, tmp_path)


def test_service_opening_songs_are_independent_complete_resources():
    root = Path(__file__).resolve().parents[1]
    specs = [
        line.split("=", 1)[1]
        for line in (root / "services/2026-10-25.flags").read_text().splitlines()
        if line.startswith("--hymns=")
    ]
    first, second = [search_hymn_md(spec, root / "processed")[0] for spec in specs[:2]]
    assert first.title == "祢的信實廣大"
    assert first.page_order == [1, 2, 3, 4, 5, 6, 6]
    assert second.title == "主我願像祢"
    assert second.page_order == [2, 1, 2, 3, 4, 9, 10]
    assert len(second.lyrics) == 10
